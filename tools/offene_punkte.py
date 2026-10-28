#!/usr/bin/env python3
"""Pflegt die Liste der offenen Punkte: schreibt OFFENE_PUNKTE.md und das Blatt „Offene Punkte“
im Datenfeld-Verzeichnis. Punkte hier bearbeiten, dann:  python3 tools/offene_punkte.py"""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT = Path(__file__).resolve().parent.parent
STAND = "28.10.2026"
# (Bereich, Offener Punkt, Wo, Wer, Priorität)
PUNKTE = [
    # --- Fachlich ---
    ("Fachlich", "Regel für genau 1× Ja im telefonischen Aspirationscheck festlegen (Vorschlag: mittleres Risiko, FEES-Entscheidung nach KSU mit Corinna)", "DYSPH-S1, SOP FEES 6.1", "Corinna, Max", "vor Go-live"),
    ("Fachlich", "Fachliche Freigabe aller Datenfeld-Verzeichnisse (alle Module stehen auf ‚Entwurf zur fachlichen Freigabe‘)", "Excel, alle Module", "Corinna", "vor Go-live"),
    ("Fachlich", "BIAS A&R: Untertest-Zuordnung, Maximalwerte und Einstufungsgrenzen mit dem Durchführungsbuch abgleichen", "BI-07 bis BI-19", "Fachliche Leitung", "vor Go-live"),
    ("Fachlich", "BoDyS: Formel für den Gesamtschweregrad hinterlegen (Verlagssoftware nicht verfügbar) und Kürzel der Stimmskalen vom Bogen übernehmen", "BD-06 bis BD-08, BD-14", "Max", "vor Go-live"),
    ("Fachlich", "KommPaS: klären, ab wann nutzbar; bis dahin nur Einschätzung VT-01", "VT-02, VT-03, DYSAR-S2", "Fachliche Leitung", "später"),
    ("Fachlich", "Kognitive Kommunikationsstörung: Verlaufsinstrument festlegen (KKS-O1 offen); MoCA nur mit Zertifizierung", "Screening KKS", "Fachliche Leitung", "später"),
    ("Fachlich", "Module ohne spezifisches Assessment: Sprechapraxie (HWL/DIAS), kognitive Kommunikationsstörung, Fazialisparese, Kinder (SP1–SP3, Kinder-Dysphagie)", "MV-06, Screening", "Fachliche Leitung", "später"),
    ("Fachlich", "Dysphagie hat keinen Block ‚Zwischenbewertung‘ wie die anderen Module (GAS, Wiederholung GUSS/FEES, Entscheidung)", "Modul DYSPH", "Fachliche Leitung", "später"),
    ("Fachlich", "Red Flags je Störungsbild in den allgemeinen Teil übernehmen (dann entfallen MA-13 und DP-05)", "MA-13, DP-05", "Max, Corinna", "später"),
    ("Fachlich", "Aktualität des ärztlichen Kehlkopfbefunds definieren (wie alt darf DP-02 sein?)", "DP-02", "Fachliche Leitung", "später"),
    ("Fachlich", "Freigabe-Regel bei FEES (Teambesprechung ab PAS 5, Gegenzeichnung bei FEDSS 5–6/Strukturbefund) in Langdock abbilden", "FE-53, Freigabe-Knoten", "Corinna, Max", "vor Go-live"),
    # --- Lizenzen ---
    ("Lizenzen", "Testmaterial und Lizenzen beschaffen bzw. prüfen: BIAS (NAT-Verlag), BoDyS (Hogrefe), CPIB-10, CETI, VHI-12, SAQOL-39g, EAT-10, GUSS-Bogen; MoCA-Zertifizierung", "Screening-Module", "Max", "vor Go-live"),
    ("Lizenzen", "Keine 1:1-Digitalisierung der Protokollbögen ohne Zustimmung der Verlage (BIAS, BoDyS) – nur Ergebniswerte erfassen", "BI-, BD-Felder", "Max", "laufend"),
    # --- tinana ---
    ("tinana", "Sandbox-Zugang: Endpunkte, Feldnamen und Authentifizierung für Patient:in und Verordnung (READ) eintragen; danach TESTMODUS = false", "langdock/2_code_tinana_lesen.js (MAPPING)", "tinana, Max", "vor Go-live"),
    ("tinana", "WRITE: Endpunkt, Feldnamen und Ablagekategorien (Befunde, Arztkorrespondenz) für die PDFs", "HTTP-Knoten tinana_befund / tinana_arzt", "tinana, Max", "vor Go-live"),
    ("tinana", "Rohdaten (JSON) zusätzlich zu PDFs übermitteln – Voraussetzung für Verläufe (FEES später, Zwischenbewertungen, Messzeitpunkte)", "Export ‚felder‘", "tinana", "später"),
    ("tinana", "Versand der Arztrückmeldung per KIM aus tinana (Backoffice, manuell) als Ablauf festlegen", "Benachrichtigung backoffice", "Max", "vor Go-live"),
    # --- Langdock ---
    ("Langdock", "Workflow in Langdock anlegen und mit allen sechs Testfällen durchlaufen (lokal getestet, in Langdock noch nicht)", "langdock/testfaelle", "Max", "vor Go-live"),
    ("Langdock", "Agent-Texte fachlich prüfen – je Modul ein Testlauf, insbesondere: erfindet der Agent nichts, stimmen Scores und Empfehlungen?", "3_agent_anweisung.md", "Corinna", "vor Go-live"),
    ("Langdock", "Webhook aus dem Browser (CORS) klären – sonst bleibt der Umweg mit Secret in der URL", "Formular ‚An Langdock senden‘", "Langdock", "vor Go-live"),
    ("Langdock", "Aufbewahrungsdauer der Workflow-Läufe (enthalten Gesundheitsdaten) klären und kürzen", "Workflow-Einstellungen", "Langdock, Max", "vor Go-live"),
    ("Langdock", "Freigaben gezielt an Personen senden (z. B. Fachliche Leitung nur bei FEES) – Empfängerliste gilt sonst für den ganzen Workflow", "Freigabe-Knoten", "Langdock", "später"),
    ("Langdock", "tinana-API-Schlüssel sicher hinterlegen statt im Code", "2_code_tinana_lesen.js", "Langdock, Max", "vor Go-live"),
    # --- Datenschutz / Recht ---
    ("Datenschutz", "AV-Verträge mit Langdock und tinana inkl. Verschwiegenheitsverpflichtung (§ 203 StGB)", "–", "Max, Anwalt Michael", "vor Go-live"),
    ("Datenschutz", "Verarbeitungsverzeichnis, TOMs und Datenschutz-Folgenabschätzung um Formular und Workflow ergänzen", "–", "Max, DSB", "vor Go-live"),
    ("Datenschutz", "Einwilligungstexte für Audio-/Video-/Fotoaufnahmen an das neue Feld MV-49 anpassen (Vertragspaket beim Anwalt)", "MV-49", "Anwalt Michael", "vor Go-live"),
    # --- Betrieb ---
    ("Betrieb", "Formular nicht über öffentliches GitHub Pages mit echten Daten nutzen: per Jamf als lokale Datei auf die Raum-Laptops verteilen oder hinter Login hosten", "index.html", "Max, netgeo", "vor Go-live"),
    ("Betrieb", "Langdock-Verbindung (Webhook-URL und Secret) auf jedem Raum-Laptop einrichten", "Formular ⚙", "Max", "vor Go-live"),
    ("Betrieb", "Terminart für FEES an einem eigenen Termin ergänzen (bisher: Ersttermin + ‚JSON laden‘)", "VO-25", "Max", "später"),
    # --- Branding ---
    ("Branding", "Quicksand als CI-Schrift bestätigen, Fußzeile und Praxisangaben im PDF prüfen", "tools/pdf_node_template.js", "Max, Rilana", "vor Go-live"),
    ("Branding", "Unterschrift im PDF mit vollem Namen statt Kürzel (Zuordnung Kürzel → Name)", "VO-24, PDF", "Max", "später"),
    # --- Dokumentation ---
    ("Dokumentation", "Älteres Google-/Claude-Dokument ‚HYO Health – Anamnesebogen‘ und der Logikbaum vom 23.10. sind veraltet – Excel ist die Quelle", "Doc vom 23.10.2026", "Max", "später"),
]

HDR = ["Nr.", "Bereich", "Offener Punkt", "Wo", "Wer", "Priorität", "Status"]


def main():
    # Markdown
    md = [f"# Offene Punkte – HYO Anamnese & Erstbericht", "", f"Stand: {STAND}. Gepflegt in `tools/offene_punkte.py`; dieselbe Liste steht im Excel (Blatt „Offene Punkte“).", ""]
    for prio in ("vor Go-live", "laufend", "später"):
        items = [(i, p) for i, p in enumerate(PUNKTE, 1) if p[4] == prio]
        if not items:
            continue
        md += [f"## {prio[0].upper() + prio[1:]} ({len(items)})", "", "| Nr. | Bereich | Offener Punkt | Wo | Wer |", "| --- | --- | --- | --- | --- |"]
        md += [f"| {i} | {p[0]} | {p[1]} | {p[2]} | {p[3]} |" for i, p in items]
        md.append("")
    (ROOT / "OFFENE_PUNKTE.md").write_text("\n".join(md), encoding="utf-8")
    # Excel
    p = ROOT / "data" / "HYO_Anamnese_Datenfelder.xlsx"
    wb = load_workbook(p)
    if "Offene Punkte" in wb.sheetnames:
        del wb["Offene Punkte"]
    ws = wb.create_sheet("Offene Punkte", index=1)
    thin = Side(style="thin", color="C8D0D6"); B = Border(left=thin, right=thin, top=thin, bottom=thin)
    rot = PatternFill("solid", fgColor="FBE3E1"); gelb = PatternFill("solid", fgColor="FFF6C8")
    for j, h in enumerate(HDR, 1):
        c = ws.cell(1, j, h); c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10); c.fill = PatternFill("solid", fgColor="1F3A4D"); c.border = B
    for i, pt in enumerate(PUNKTE, 1):
        row = [i, pt[0], pt[1], pt[2], pt[3], pt[4], "offen"]
        for j, v in enumerate(row, 1):
            c = ws.cell(i + 1, j, v); c.font = Font(name="Arial", size=10); c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = B
            if j == 6:
                c.fill = rot if v == "vor Go-live" else gelb
    for col, w in zip("ABCDEFG", (6, 14, 80, 34, 20, 13, 10)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:G{len(PUNKTE) + 1}"
    wb.save(p)
    print(f"{len(PUNKTE)} offene Punkte, davon {sum(1 for x in PUNKTE if x[4] == 'vor Go-live')} vor Go-live")


if __name__ == "__main__":
    main()
