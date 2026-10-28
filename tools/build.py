#!/usr/bin/env python3
"""Erzeugt index.html aus dem Datenfeld-Verzeichnis (Excel) und dem HTML-Template.

Aufruf (im Repo-Hauptordner):
    python3 tools/build.py

Liest  data/HYO_Anamnese_Datenfelder.xlsx und tools/template.html,
schreibt index.html. Bricht ab, wenn eine Bedingung nicht gelesen werden kann
oder eine Werteliste fehlt.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
XLSX = ROOT / "data" / "HYO_Anamnese_Datenfelder.xlsx"
TEMPLATE = ROOT / "tools" / "template.html"
OUT = ROOT / "index.html"

ID = r"[A-Z]{2}-\d{2}"


def parse_atom(text):
    t = text.strip()
    m = re.fullmatch(rf"({ID}) enthält (.+)", t)
    if m:
        vals = [v.strip() for v in re.split(r"\s*,\s*|\s+oder\s+", m.group(2)) if v.strip()]
        return {"op": "has", "f": m.group(1), "v": vals}
    m = re.fullmatch(rf"({ID}) ∈ (.+)", t)
    if m:
        return {"op": "in", "f": m.group(1), "v": [v.strip() for v in re.split(r"\s*,\s*|\s+oder\s+", m.group(2)) if v.strip()]}
    m = re.fullmatch(r"Alter (<|≥) (\d+)", t)
    if m:  # Alter in Jahren aus VO-02 (Geburtsdatum) und VO-23 (Aufnahmedatum, sonst heute)
        return {"op": "age_lt" if m.group(1) == "<" else "age_ge", "f": "VO-02", "v": int(m.group(2))}
    m = re.fullmatch(rf"({ID}) (=|≠|≥|≤) (.+)", t)
    if m:
        op = {"=": "eq", "≠": "ne", "≥": "ge", "≤": "le"}[m.group(2)]
        return {"op": op, "f": m.group(1), "v": m.group(3).strip()}
    m = re.fullmatch(rf"({ID}) leer", t)
    if m:
        return {"op": "isempty", "f": m.group(1)}
    m = re.fullmatch(rf"({ID}) (angezeigt|befüllt|ausgefüllt)", t)
    if m:
        return {"op": "shown" if m.group(2) == "angezeigt" else "filled", "f": m.group(1)}
    m = re.fullmatch(r"Eintrag (\d) befüllt", t)
    if m:  # zusätzliche Diagnosegruppe n -> Feld MV-33 / MV-36 / MV-39
        return {"op": "filled", "f": f"MV-{30 + 3 * int(m.group(1)):02d}"}
    raise ValueError(f"Unbekannte Bedingung: {text!r}")


def parse_condition(text):
    if not text:
        return None
    if text.strip() == "je Code in MV-06":
        return None  # Screening-Block wird je Modul gerendert
    ors = re.split(r"\s+(?:ODER|oder)\s+(?=[A-Z]{2}-\d)", text.strip())
    out = []
    for part in ors:
        ands = re.split(r"\s+(?:UND|und)\s+(?=[A-Z]{2}-\d|Eintrag|Alter)", part)
        out.append([parse_atom(a) for a in ands])
    return out  # Liste von UND-Gruppen, verknüpft mit ODER


def rows(ws):
    hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in r):
            yield dict(zip(hdr, r))


def main():
    wb = load_workbook(XLSX, data_only=True)

    lists = {}
    for r in rows(wb["Wertelisten"]):
        lists.setdefault(r["Liste"], []).append({"c": str(r["Code"]), "l": r["Bezeichnung"]})

    fields = []
    for r in rows(wb["Datenfelder"]):
        cond_text = r.get("Bedingung (anzeigen/Pflicht wenn)")
        src = r.get("Quelle") or ""
        if src.startswith("berechnet"):
            src = "berechnet"  # Details der Berechnung stehen im Hinweis bzw. in der Legende
        lst = r.get("Werteliste")
        if lst and lst not in lists:
            sys.exit(f"Werteliste fehlt: {lst} ({r['ID']})")
        fields.append({
            "id": r["ID"], "key": r["Schlüssel"], "sec": r["Abschnitt"], "grp": r["Gruppe"],
            "label": r["Feldname"], "type": r["Datentyp"], "fmt": r.get("Format / Validierung") or "",
            "list": lst, "req": r["Pflicht"], "src": src,
            "condText": cond_text or "", "cond": parse_condition(cond_text),
            "ex": "" if r.get("Beispielwert") is None else str(r["Beispielwert"]),
            "hint": r.get("Hinweis") or "", "scope": r.get("Gilt für") or "alle",
            "mod": r.get("Modul") or "",
            "excl": (re.search(r"\b([A-Z_]+) schließt andere Codes aus", r.get("Hinweis") or "") or [None, None])[1],
        })

    modules = []
    for r in rows(wb["Screening-Module"]):
        modules.append({"mod": r["Modul-Code"], "vd": r["Verdachtsdiagnose"], "dg": r["Typ. Diagnosegruppe"],
                        "stage": r["Stufe"], "iid": r["Instrument-ID"], "name": r["Instrument"],
                        "fmt": r["Ergebnisformat"] or ""})

    guss = []
    ws = wb["GUSS-Auswertung"]
    for r in ws.iter_rows(min_row=2, max_col=6, values_only=True):
        if r[0] and r[1]:
            guss.append({"range": str(r[0]), "befund": r[1], "grad": r[2], "sp": r[3], "ge": r[4], "more": r[5]})

    ablauf = [list(r) for r in wb["Ablauf Ersttermin SC"].iter_rows(min_row=2, values_only=True) if r[0]]

    def kost(txt):  # "4 – püriert" -> "4", "NPO (nichts oral)" -> "NPO"
        t = str(txt or "").strip()
        return "NPO" if t.startswith("NPO") else t.split(" ")[0]
    fedss = {}
    if "FEDSS-Auswertung" in wb.sheetnames:
        for r in wb["FEDSS-Auswertung"].iter_rows(min_row=2, max_col=6, values_only=True):
            if r[0] is not None and str(r[0]).strip().isdigit():
                fedss[str(r[0]).strip()] = {"befund": r[1], "empf": r[2], "sp": kost(r[3]), "ge": kost(r[4]),
                                            "sp_txt": r[3], "ge_txt": r[4], "more": r[5] if r[5] not in (None, "–") else ""}

    ids = [f["id"] for f in fields]
    # Prüfung: Codes in Bedingungen müssen in der Werteliste des Felds vorkommen
    by = {f["id"]: f for f in fields}
    fehler = []
    for f in fields:
        for grp in f["cond"] or []:
            for a in grp:
                ref = by.get(a["f"])
                if not ref or a["op"] not in ("eq", "ne", "has", "in"):
                    continue
                codes = {c["c"].lower() for c in lists.get(ref["list"], [])} if ref["list"] else ({"ja", "nein"} if ref["type"] == "boolean" else None)
                if codes is None:
                    continue
                for v in (a["v"] if isinstance(a["v"], list) else [a["v"]]):
                    if str(v).lower() not in codes:
                        fehler.append(f"{f['id']}: Code {v!r} gibt es in {ref['list'] or 'ja/nein'} ({a['f']}) nicht")
    for f in fields:  # Beispielwerte gegen Wertelisten
        if f["list"] and f["ex"] and f["type"].startswith("enum"):
            codes = {c["c"] for c in lists[f["list"]]}
            for v in [x.strip() for x in f["ex"].split(";")]:
                if v not in codes:
                    fehler.append(f"{f['id']}: Beispielwert {v!r} nicht in {f['list']}")
    if fehler:
        sys.exit("Fehler im Datenfeld-Verzeichnis:\n  " + "\n  ".join(fehler))
    for f in fields:
        for grp in f["cond"] or []:
            for a in grp:
                if a["f"] not in ids and a["f"] != "SC-06":
                    sys.exit(f"{f['id']}: Bedingung verweist auf unbekanntes Feld {a['f']}")

    data = {"fields": fields, "lists": lists, "modules": modules, "guss": guss, "ablauf": ablauf, "fedss": fedss,
            "source": XLSX.name}
    payload = json.dumps(data, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")
    html = TEMPLATE.read_text(encoding="utf-8").replace("/*__DATA__*/null", payload)
    OUT.write_text(html, encoding="utf-8")
    print(f"index.html geschrieben: {len(fields)} Felder, {len(lists)} Wertelisten, {len(modules)} Instrumente")


if __name__ == "__main__":
    main()
